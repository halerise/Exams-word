from flask import Flask, render_template, request, redirect, url_for, session
import sqlite3
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO
from flask import send_file

app = Flask(__name__)
app.secret_key = "your_secret_key"  # Change this to a random secret key

ADMIN_PASSWORD = "admin@12345678"

DATABASE = "database.db"


# -----------------------------
# DATABASE
# -----------------------------

def get_db():
    conn = sqlite3.connect(DATABASE)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    conn = get_db()

    conn.execute("""
        CREATE TABLE IF NOT EXISTS results (
            id INTEGER PRIMARY KEY AUTOINCREMENT,

            student1_id TEXT NOT NULL,
            student1_name TEXT NOT NULL,

            student2_id TEXT NOT NULL,
            student2_name TEXT NOT NULL,

            score INTEGER NOT NULL,
            total INTEGER NOT NULL,
            percentage REAL NOT NULL,
            grade TEXT NOT NULL,

            submitted_at TEXT NOT NULL
        )
    """)

    conn.commit()
    conn.close()

  


# -----------------------------
# QUESTIONS
# -----------------------------

questions = [
    {
        "id": 1,
        "question": "What is a computer?",
        "options": [
            "An electronic device that processes data",
            "A type of television",
            "A type of telephone",
            "A paper document"
        ],
        "answer": 0
    },
    {
        "id": 2,
        "question": "Which of the following is an input device?",
        "options": [
            "Monitor",
            "Printer",
            "Keyboard",
            "Speaker"
        ],
        "answer": 2
    },
    {
        "id": 3,
        "question": "Which component is often called the brain of the computer?",
        "options": [
            "Keyboard",
            "CPU",
            "Monitor",
            "Mouse"
        ],
        "answer": 1
    },
    {
        "id": 4,
        "question": "Which of the following is used to store files permanently?",
        "options": [
            "Hard drive/SSD",
            "Keyboard",
            "Monitor",
            "Mouse"
        ],
        "answer": 0
    },
    {
        "id": 5,
        "question": "What does RAM stand for?",
        "options": [
            "Read Access Memory",
            "Random Access Memory",
            "Run Access Machine",
            "Random Application Memory"
        ],
        "answer": 1
    },
    {
        "id": 6,
        "question": "Which of these is an operating system?",
        "options": [
            "Microsoft Word",
            "Windows",
            "Google Chrome",
            "Microsoft Excel"
        ],
        "answer": 1
    },
    {
        "id": 7,
        "question": "Which device is mainly used to point, click and select items?",
        "options": [
            "Mouse",
            "Printer",
            "Speaker",
            "Scanner"
        ],
        "answer": 0
    },
    {
        "id": 8,
        "question": "Which of the following is an output device?",
        "options": [
            "Keyboard",
            "Mouse",
            "Monitor",
            "Scanner"
        ],
        "answer": 2
    },
    {
        "id": 9,
        "question": "Which one is an example of computer software?",
        "options": [
            "Keyboard",
            "Monitor",
            "Microsoft Word",
            "Hard disk"
        ],
        "answer": 2
    },
    {
        "id": 10,
        "question": "Which keyboard shortcut is commonly used to copy selected text?",
        "options": [
            "Ctrl + X",
            "Ctrl + C",
            "Ctrl + V",
            "Ctrl + Z"
        ],
        "answer": 1
    },

    # MICROSOFT WORD

    {
        "id": 11,
        "question": "What is Microsoft Word mainly used for?",
        "options": [
            "Creating and editing documents",
            "Browsing the internet",
            "Playing music",
            "Managing computer hardware"
        ],
        "answer": 0
    },
    {
        "id": 12,
        "question": "Which file extension is commonly associated with modern Microsoft Word documents?",
        "options": [
            ".xlsx",
            ".pptx",
            ".docx",
            ".jpg"
        ],
        "answer": 2
    },
    {
        "id": 13,
        "question": "Which shortcut is used to save a Word document?",
        "options": [
            "Ctrl + S",
            "Ctrl + P",
            "Ctrl + O",
            "Ctrl + N"
        ],
        "answer": 0
    },
    {
        "id": 14,
        "question": "Which shortcut creates a new document in Microsoft Word?",
        "options": [
            "Ctrl + N",
            "Ctrl + B",
            "Ctrl + F",
            "Ctrl + H"
        ],
        "answer": 0
    },
    {
        "id": 15,
        "question": "Which shortcut is used to open an existing document?",
        "options": [
            "Ctrl + O",
            "Ctrl + P",
            "Ctrl + E",
            "Ctrl + L"
        ],
        "answer": 0
    },
    {
        "id": 16,
        "question": "What does Ctrl + Z normally do in Microsoft Word?",
        "options": [
            "Saves the document",
            "Prints the document",
            "Undoes the previous action",
            "Closes Word"
        ],
        "answer": 2
    },
    {
        "id": 17,
        "question": "Which shortcut is used to make selected text bold?",
        "options": [
            "Ctrl + I",
            "Ctrl + U",
            "Ctrl + B",
            "Ctrl + D"
        ],
        "answer": 2
    },
    {
        "id": 18,
        "question": "Which shortcut is used to italicize selected text?",
        "options": [
            "Ctrl + I",
            "Ctrl + B",
            "Ctrl + U",
            "Ctrl + T"
        ],
        "answer": 0
    },
    {
        "id": 19,
        "question": "Which shortcut is used to underline selected text?",
        "options": [
            "Ctrl + B",
            "Ctrl + U",
            "Ctrl + I",
            "Ctrl + L"
        ],
        "answer": 1
    },
    {
        "id": 20,
        "question": "Which alignment places text in the center of the page?",
        "options": [
            "Left",
            "Right",
            "Justify",
            "Center"
        ],
        "answer": 3
    },
    {
        "id": 21,
        "question": "Which alignment makes text line up evenly along both the left and right margins?",
        "options": [
            "Left",
            "Center",
            "Justify",
            "Right"
        ],
        "answer": 2
    },
    {
        "id": 22,
        "question": "Which Word feature is used to create a list with dots or other symbols?",
        "options": [
            "Bullets",
            "Margins",
            "Columns",
            "WordArt"
        ],
        "answer": 0
    },
    {
        "id": 23,
        "question": "Which feature is used to organize information into rows and columns?",
        "options": [
            "Table",
            "Header",
            "Footer",
            "Page Break"
        ],
        "answer": 0
    },
    {
        "id": 24,
        "question": "What is the purpose of margins in a Word document?",
        "options": [
            "They control the blank space around the edges of the page",
            "They change the font",
            "They insert pictures",
            "They check spelling"
        ],
        "answer": 0
    },
    {
        "id": 25,
        "question": "Which option changes a page from Portrait to Landscape?",
        "options": [
            "Size",
            "Orientation",
            "Margins",
            "Columns"
        ],
        "answer": 1
    },
    {
        "id": 26,
        "question": "Where can you normally find information such as a document title at the top of every page?",
        "options": [
            "Footer",
            "Header",
            "Status bar",
            "Ribbon"
        ],
        "answer": 1
    },
    {
        "id": 27,
        "question": "Which Word feature allows you to search for a particular word in a document?",
        "options": [
            "Find",
            "Save",
            "Print",
            "Zoom"
        ],
        "answer": 0
    },
    {
        "id": 28,
        "question": "Which feature allows you to replace one word with another throughout a document?",
        "options": [
            "Word Count",
            "Replace",
            "Translate",
            "Zoom"
        ],
        "answer": 1
    },
    {
        "id": 29,
        "question": "Which shortcut opens the Print window in Microsoft Word?",
        "options": [
            "Ctrl + P",
            "Ctrl + R",
            "Ctrl + W",
            "Ctrl + G"
        ],
        "answer": 0
    },
    {
        "id": 30,
        "question": "Which feature checks a document for spelling and grammar errors?",
        "options": [
            "Mail Merge",
            "Spelling & Grammar",
            "Page Color",
            "WordArt"
        ],
        "answer": 1
    }
]


# -----------------------------
# GRADING
# -----------------------------

def calculate_grade(percentage):

    if percentage >= 80:
        return "A"
    elif percentage >= 70:
        return "B"
    elif percentage >= 60:
        return "C"
    elif percentage >= 50:
        return "D"
    else:
        return "E"


# -----------------------------
# LOGIN
# -----------------------------
@app.route("/", methods=["GET", "POST"])
def login():

    if request.method == "POST":

        student1_id = request.form.get("student1_id")
        student1_name = request.form.get("student1_name")

        student2_id = request.form.get("student2_id")
        student2_name = request.form.get("student2_name")

        if not student1_id or not student1_name:
            return "Please enter Student 1 details."

        if not student2_id or not student2_name:
            return "Please enter Student 2 details."

        return redirect(
            url_for(
                "exam",
                student1_id=student1_id,
                student1_name=student1_name,
                student2_id=student2_id,
                student2_name=student2_name
            )
        )

    return render_template("login.html")



# -----------------------------
# EXAM
# -----------------------------

@app.route("/exam", methods=["GET", "POST"])
def exam():

    if request.method == "GET":

        student1_id = request.args.get("student1_id")
        student1_name = request.args.get("student1_name")

        student2_id = request.args.get("student2_id")
        student2_name = request.args.get("student2_name")

        return render_template(
            "exam.html",
            questions=questions,

            student1_id=student1_id,
            student1_name=student1_name,

            student2_id=student2_id,
            student2_name=student2_name
        )

    # -------------------------
    # GET STUDENT INFORMATION
    # -------------------------

    student1_id = request.form.get("student1_id")
    student1_name = request.form.get("student1_name")

    student2_id = request.form.get("student2_id")
    student2_name = request.form.get("student2_name")

    # Make sure all student information exists

    if not student1_id or not student1_name:
        return "Student 1 information is missing."

    if not student2_id or not student2_name:
        return "Student 2 information is missing."


    # -------------------------
    # PREVENT DUPLICATE
    # -------------------------

    conn = get_db()

    existing_result = conn.execute("""
        SELECT id
        FROM results
        WHERE student1_id = ?
           OR student2_id = ?
           OR student1_id = ?
           OR student2_id = ?
        LIMIT 1
    """, (
        student1_id,
        student1_id,
        student2_id,
        student2_id
    )).fetchone()

    conn.close()

    if existing_result:

        return """
        <h2>Exam Already Submitted</h2>

        <p>
            One or both of these student IDs have already
            submitted an examination.
        </p>

        <a href="/">Return to Login</a>
        """


    # -------------------------
    # MARK EXAM
    # -------------------------

    score = 0

    answers = {}

    for question in questions:

        question_id = str(question["id"])

        selected_answer = request.form.get(
            f"question_{question_id}"
        )

        answers[question_id] = selected_answer

        if selected_answer is not None:

            if int(selected_answer) == question["answer"]:
                score += 1


    # -------------------------
    # CALCULATE RESULT
    # -------------------------

    total = len(questions)

    percentage = (score / total) * 100

    grade = calculate_grade(percentage)

    submitted_at = datetime.now().strftime(
        "%Y-%m-%d %H:%M:%S"
    )


    # -------------------------
    # SAVE RESULT
    # -------------------------

    conn = get_db()

    conn.execute("""
        INSERT INTO results
        (
            student1_id,
            student1_name,
            student2_id,
            student2_name,
            score,
            total,
            percentage,
            grade,
            submitted_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        student1_id,
        student1_name,
        student2_id,
        student2_name,
        score,
        total,
        percentage,
        grade,
        submitted_at
    ))

    conn.commit()
    conn.close()


    # -------------------------
    # SHOW RESULT
    # -------------------------

    return render_template(
        "result.html",

        questions=questions,
        answers=answers,

        score=score,
        total=total,
        percentage=percentage,
        grade=grade,

        student1_name=student1_name,
        student2_name=student2_name
    )
    # -------------------------
    # EXAM SUBMISSION
    # -------------------------

    student1_id = request.form.get("student1_id")
    student1_name = request.form.get("student1_name")

    student2_id = request.form.get("student2_id")
    student2_name = request.form.get("student2_name")

    score = 0

    answers = {}

    for question in questions:

        question_id = str(question["id"])

        selected_answer = request.form.get(
            f"question_{question_id}"
        )

        answers[question_id] = selected_answer

        if selected_answer is not None:

            if int(selected_answer) == question["answer"]:
                score += 1

    total = len(questions)

    percentage = (score / total) * 100

    grade = calculate_grade(percentage)

    submitted_at = datetime.now().strftime(
        "%Y-%m-%d %H:%M:%S"
    )

    conn = get_db()

    conn.execute("""
        INSERT INTO results
        (
            student1_id,
            student1_name,
            student2_id,
            student2_name,
            score,
            total,
            percentage,
            grade,
            submitted_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        student1_id,
        student1_name,
        student2_id,
        student2_name,
        score,
        total,
        percentage,
        grade,
        submitted_at
    ))

    conn.commit()
    conn.close()

    return render_template(
        "result.html",

        questions=questions,
        answers=answers,

        score=score,
        total=total,
        percentage=percentage,
        grade=grade,

        student1_name=student1_name,
        student2_name=student2_name
    )




@app.route("/admin-login", methods=["GET", "POST"])
def admin_login():

    if request.method == "POST":

        password = request.form.get("password")

        if password == ADMIN_PASSWORD:

            session["admin_logged_in"] = True

            return redirect(url_for("admin"))

        return render_template(
            "admin_login.html",
            error="Incorrect password."
        )

    return render_template("admin_login.html")
# -----------------------------
# ADMIN RESULTS
# -----------------------------

@app.route("/admin")
def admin():

    if not session.get("admin_logged_in"):
        return redirect(url_for("admin_login"))

    conn = get_db()

    results = conn.execute("""
        SELECT *
        FROM results
        ORDER BY score DESC, submitted_at ASC
    """).fetchall()

    conn.close()

    ranked_results = []

    previous_score = None
    current_rank = 0

    for index, result in enumerate(results):

        if result["score"] != previous_score:

            current_rank = index + 1
            previous_score = result["score"]

        ranked_results.append({
            "rank": current_rank,

            "student1_id": result["student1_id"],
            "student1_name": result["student1_name"],

            "student2_id": result["student2_id"],
            "student2_name": result["student2_name"],

            "score": result["score"],
            "total": result["total"],
            "percentage": result["percentage"],
            "grade": result["grade"],
            "submitted_at": result["submitted_at"]
        })

    total_students = len(ranked_results)

    average_score = 0

    if total_students > 0:

        average_score = sum(
            result["percentage"]
            for result in ranked_results
        ) / total_students

    return render_template(
        "admin.html",
        results=ranked_results,
        total_students=total_students,
        average_score=average_score
    )


@app.route("/admin-logout")
def admin_logout():

    session.pop("admin_logged_in", None)

    return redirect(url_for("admin_login"))

@app.route("/export-results")
def export_results():

    # Protect the export
    if not session.get("admin_logged_in"):
        return redirect(url_for("admin_login"))

    conn = get_db()

    results = conn.execute("""
        SELECT *
        FROM results
        ORDER BY score DESC, submitted_at ASC
    """).fetchall()

    conn.close()

    # Create workbook
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "CAT Results"

    # Title
    worksheet.merge_cells("A1:I1")

    worksheet["A1"] = "COMPUTER APPLICATIONS CAT RESULTS"

    worksheet["A1"].font = Font(
        bold=True,
        size=16
    )

    worksheet["A1"].alignment = Alignment(
        horizontal="center"
    )

    # Subtitle
    worksheet.merge_cells("A2:I2")

    worksheet["A2"] = "Microsoft Word & Basic Computer Knowledge"

    worksheet["A2"].font = Font(
        italic=True,
        size=11
    )

    worksheet["A2"].alignment = Alignment(
        horizontal="center"
    )

    # Headers
    headers = [
        "Position",
        "Student 1 ID",
        "Student 1 Name",
        "Student 2 ID",
        "Student 2 Name",
        "Score",
        "Percentage",
        "Grade",
        "Submitted At"
    ]

    header_row = 4

    for column, header in enumerate(headers, start=1):

        cell = worksheet.cell(
            row=header_row,
            column=column,
            value=header
        )

        cell.font = Font(bold=True)

        cell.alignment = Alignment(
            horizontal="center"
        )

    # Ranking
    previous_score = None
    current_rank = 0

    # Data starts on row 5
    for index, result in enumerate(results):

        score = result["score"]

        if score != previous_score:
            current_rank = index + 1
            previous_score = score

        row = index + 5

        data = [
            current_rank,
            result["student1_id"],
            result["student1_name"],
            result["student2_id"],
            result["student2_name"],
            f"{result['score']}/{result['total']}",
            result["percentage"] / 100,
            result["grade"],
            result["submitted_at"]
        ]

        for column, value in enumerate(data, start=1):

            cell = worksheet.cell(
                row=row,
                column=column,
                value=value
            )

            cell.alignment = Alignment(
                vertical="center"
            )

        # Format percentage
        worksheet.cell(
            row=row,
            column=7
        ).number_format = "0.0%"

    # Borders
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    for row in worksheet.iter_rows(
        min_row=4,
        max_row=worksheet.max_row,
        min_col=1,
        max_col=9
    ):

        for cell in row:
            cell.border = thin_border

    # Column widths
    widths = {
        "A": 12,
        "B": 18,
        "C": 25,
        "D": 18,
        "E": 25,
        "F": 12,
        "G": 15,
        "H": 10,
        "I": 22
    }

    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width

    # Freeze headers
    worksheet.freeze_panes = "A5"

    # Create file in memory
    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="Computer_CAT_Results.xlsx",
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )



# -----------------------------
# START APPLICATION
# -----------------------------

if __name__ == "__main__":

    init_db()

    app.run(
        host="0.0.0.0",
        port=5000,
        debug=True
    )

