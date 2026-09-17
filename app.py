from flask import Flask, render_template, request, redirect, url_for, session
import sqlite3
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO
from flask import send_file

app = Flask(__name__)
app.secret_key = "your_secret_key"  # Change this to a random secret key

ADMIN_PASSWORD = "admin@12345678"

DATABASE = "database.db"


# -----------------------------
# DATABASE
# -----------------------------

def get_db():
    conn = sqlite3.connect(DATABASE)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    conn = get_db()

    conn.execute("""
        CREATE TABLE IF NOT EXISTS results (
            id INTEGER PRIMARY KEY AUTOINCREMENT,

            student1_id TEXT NOT NULL,
            student1_name TEXT NOT NULL,

            student2_id TEXT NOT NULL,
            student2_name TEXT NOT NULL,

            score INTEGER NOT NULL,
            total INTEGER NOT NULL,
            percentage REAL NOT NULL,
            grade TEXT NOT NULL,

            submitted_at TEXT NOT NULL
        )
    """)

    conn.commit()
    conn.close()

  


# -----------------------------
# ADVANCED CAT QUESTIONS
# -----------------------------

questions = [

    # =============================
    # MICROSOFT WORD
    # =============================

    {
        "id": 1,
        "question": "A document contains several chapters. Which Word feature is most appropriate for creating an automatic Table of Contents?",
        "options": [
            "WordArt",
            "Heading Styles",
            "Text Boxes",
            "Format Painter"
        ],
        "answer": 1
    },

    {
        "id": 2,
        "question": "You want the first page of a document to have no header while all other pages have a header. Which option should you use?",
        "options": [
            "Different Odd & Even Pages",
            "Different First Page",
            "Page Break",
            "Continuous Section Break"
        ],
        "answer": 1
    },

    {
        "id": 3,
        "question": "What is the main purpose of a Section Break in Microsoft Word?",
        "options": [
            "To delete a page",
            "To divide a document into independently formatted sections",
            "To correct spelling mistakes",
            "To save the document"
        ],
        "answer": 1
    },

    {
        "id": 4,
        "question": "You need pages 1–3 in Portrait orientation and pages 4–5 in Landscape orientation. What should you insert before page 4?",
        "options": [
            "Page Break",
            "Line Break",
            "Section Break",
            "Column Break"
        ],
        "answer": 2
    },

    {
        "id": 5,
        "question": "A 30-page document contains the phrase 'Computer Applications' 50 times. You want to change every occurrence to 'Computer Applications CAT'. Which feature should you use?",
        "options": [
            "Word Count",
            "Find and Replace",
            "Spelling & Grammar",
            "Format Painter"
        ],
        "answer": 1
    },

    {
        "id": 6,
        "question": "A teacher wants to see exactly which parts of a document a student has added, deleted, or modified. Which feature should be enabled?",
        "options": [
            "Comments",
            "Track Changes",
            "Word Count",
            "Navigation Pane"
        ],
        "answer": 1
    },

    {
        "id": 7,
        "question": "Which Word feature allows you to compare two versions of a document and identify their differences?",
        "options": [
            "Compare",
            "Translate",
            "Inspect Document",
            "Mail Merge"
        ],
        "answer": 0
    },

    {
        "id": 8,
        "question": "A school needs to generate 200 personalized letters using one Word document and an Excel list of student names and IDs. Which feature should be used?",
        "options": [
            "SmartArt",
            "Mail Merge",
            "WordArt",
            "AutoCorrect"
        ],
        "answer": 1
    },

    {
        "id": 9,
        "question": "In Mail Merge, an Excel spreadsheet containing student names and IDs is normally used as the:",
        "options": [
            "Main document",
            "Data source",
            "Merge field",
            "Header"
        ],
        "answer": 1
    },

    {
        "id": 10,
        "question": "A teacher wants all Heading 1 titles throughout a document to have exactly the same formatting. What is the best approach?",
        "options": [
            "Format every heading manually",
            "Use the Heading 1 Style",
            "Use WordArt",
            "Use a Text Box"
        ],
        "answer": 1
    },

    {
        "id": 11,
        "question": "If the Heading 1 style is modified, what normally happens to text that already uses Heading 1?",
        "options": [
            "Only the selected word changes",
            "The document is deleted",
            "Text using that style adopts the modified formatting",
            "All text becomes Heading 1"
        ],
        "answer": 2
    },

    {
        "id": 12,
        "question": "Which Word feature allows you to quickly navigate through a long document using its headings?",
        "options": [
            "Clipboard",
            "Navigation Pane",
            "Status Bar",
            "Ruler"
        ],
        "answer": 1
    },

    {
        "id": 13,
        "question": "You insert a photograph into a paragraph and want text to flow around the photograph. Which feature controls this?",
        "options": [
            "Wrap Text",
            "Page Color",
            "Text Direction",
            "Page Borders"
        ],
        "answer": 0
    },

    {
        "id": 14,
        "question": "Which text-wrapping option allows text to follow the approximate outline of an irregularly shaped image?",
        "options": [
            "In Line with Text",
            "Square",
            "Tight",
            "Behind Text"
        ],
        "answer": 2
    },

    {
        "id": 15,
        "question": "You have several figures in a research document and want Word to number them automatically. Which feature should you use?",
        "options": [
            "Caption",
            "Bookmark",
            "Hyperlink",
            "Footnote"
        ],
        "answer": 0
    },

    {
        "id": 16,
        "question": "What is the main purpose of a Bookmark in Microsoft Word?",
        "options": [
            "To automatically save a document",
            "To mark a specific location for quick navigation or linking",
            "To change the document font",
            "To check grammar"
        ],
        "answer": 1
    },

    {
        "id": 17,
        "question": "Which feature can be used to prevent users from modifying protected parts of a Word document?",
        "options": [
            "Read Mode",
            "Restrict Editing",
            "Word Count",
            "Format Painter"
        ],
        "answer": 1
    },

    {
        "id": 18,
        "question": "A numbered list unexpectedly starts again at 1, but you want it to continue from the previous list. Which option should you use?",
        "options": [
            "Restart Numbering",
            "Continue Numbering",
            "Set Numbering Value to 1",
            "Convert to Text"
        ],
        "answer": 1
    },

    {
        "id": 19,
        "question": "Which keyboard shortcut opens the Find and Replace dialog box in Microsoft Word?",
        "options": [
            "Ctrl + F",
            "Ctrl + H",
            "Ctrl + G",
            "Ctrl + R"
        ],
        "answer": 1
    },

    {
        "id": 20,
        "question": "Which keyboard shortcut opens the Go To function, allowing you to move directly to a specific page?",
        "options": [
            "Ctrl + G",
            "Ctrl + J",
            "Ctrl + K",
            "Ctrl + L"
        ],
        "answer": 0
    },


    # =============================
    # BASIC COMPUTER KNOWLEDGE
    # =============================

    {
        "id": 21,
        "question": "Which component temporarily stores programs and data that the CPU is actively using?",
        "options": [
            "SSD",
            "RAM",
            "ROM",
            "BIOS"
        ],
        "answer": 1
    },

    {
        "id": 22,
        "question": "Which statement best describes the role of an operating system?",
        "options": [
            "It only provides internet access",
            "It manages computer hardware and provides services for applications",
            "It permanently stores all user files",
            "It is a type of antivirus"
        ],
        "answer": 1
    },

    {
        "id": 23,
        "question": "Which of the following is an example of system software?",
        "options": [
            "Microsoft Word",
            "Microsoft Excel",
            "Windows",
            "Adobe Photoshop"
        ],
        "answer": 2
    },

    {
        "id": 24,
        "question": "What is the primary function of the CPU?",
        "options": [
            "Store files permanently",
            "Process instructions and perform calculations",
            "Display images",
            "Connect the computer to Wi-Fi"
        ],
        "answer": 1
    },

    {
        "id": 25,
        "question": "Which storage device generally provides faster data access and has no moving mechanical parts?",
        "options": [
            "HDD",
            "SSD",
            "DVD",
            "Floppy Disk"
        ],
        "answer": 1
    },

    {
        "id": 26,
        "question": "What does an IP address primarily identify in a network?",
        "options": [
            "A network interface or device",
            "The computer's RAM capacity",
            "The processor speed",
            "The monitor resolution"
        ],
        "answer": 0
    },

    {
        "id": 27,
        "question": "Which network device normally forwards packets between different networks?",
        "options": [
            "Switch",
            "Router",
            "Keyboard",
            "Monitor"
        ],
        "answer": 1
    },

    {
        "id": 28,
        "question": "Which protocol is primarily used for secure communication when browsing websites?",
        "options": [
            "HTTP",
            "FTP",
            "HTTPS",
            "SMTP"
        ],
        "answer": 2
    },

    {
        "id": 29,
        "question": "Which situation is the best example of phishing?",
        "options": [
            "Installing a legitimate Windows update",
            "Receiving a fake login message designed to steal your password",
            "Creating a Word document",
            "Compressing files into a ZIP folder"
        ],
        "answer": 1
    },

    {
        "id": 30,
        "question": "A computer can communicate with devices on its local network but cannot access devices on other networks. Which configuration should be checked first?",
        "options": [
            "Monitor resolution",
            "Default gateway",
            "Keyboard layout",
            "Screen brightness"
        ],
        "answer": 1
    }
]

# -----------------------------
# GRADING
# -----------------------------

def calculate_grade(percentage):

    if percentage >= 80:
        return "A"
    elif percentage >= 70:
        return "B"
    elif percentage >= 60:
        return "C"
    elif percentage >= 50:
        return "D"
    else:
        return "E"


# -----------------------------
# LOGIN
# -----------------------------
@app.route("/", methods=["GET", "POST"])
def login():

    if request.method == "POST":

        student1_id = request.form.get("student1_id")
        student1_name = request.form.get("student1_name")

        student2_id = request.form.get("student2_id")
        student2_name = request.form.get("student2_name")

        if not student1_id or not student1_name:
            return "Please enter Student 1 details."

        if not student2_id or not student2_name:
            return "Please enter Student 2 details."

        return redirect(
            url_for(
                "exam",
                student1_id=student1_id,
                student1_name=student1_name,
                student2_id=student2_id,
                student2_name=student2_name
            )
        )

    return render_template("login.html")



# -----------------------------
# EXAM
# -----------------------------

@app.route("/exam", methods=["GET", "POST"])
def exam():

    if request.method == "GET":

        student1_id = request.args.get("student1_id")
        student1_name = request.args.get("student1_name")

        student2_id = request.args.get("student2_id")
        student2_name = request.args.get("student2_name")

        return render_template(
            "exam.html",
            questions=questions,

            student1_id=student1_id,
            student1_name=student1_name,

            student2_id=student2_id,
            student2_name=student2_name
        )

    # -------------------------
    # GET STUDENT INFORMATION
    # -------------------------

    student1_id = request.form.get("student1_id")
    student1_name = request.form.get("student1_name")

    student2_id = request.form.get("student2_id")
    student2_name = request.form.get("student2_name")

    # Make sure all student information exists

    if not student1_id or not student1_name:
        return "Student 1 information is missing."

    if not student2_id or not student2_name:
        return "Student 2 information is missing."


    # -------------------------
    # PREVENT DUPLICATE
    # -------------------------

    conn = get_db()

    existing_result = conn.execute("""
        SELECT id
        FROM results
        WHERE student1_id = ?
           OR student2_id = ?
           OR student1_id = ?
           OR student2_id = ?
        LIMIT 1
    """, (
        student1_id,
        student1_id,
        student2_id,
        student2_id
    )).fetchone()

    conn.close()

    if existing_result:

        return """
        <h2>Exam Already Submitted</h2>

        <p>
            One or both of these student IDs have already
            submitted an examination.
        </p>

        <a href="/">Return to Login</a>
        """


    # -------------------------
    # MARK EXAM
    # -------------------------

    score = 0

    answers = {}

    for question in questions:

        question_id = str(question["id"])

        selected_answer = request.form.get(
            f"question_{question_id}"
        )

        answers[question_id] = selected_answer

        if selected_answer is not None:

            if int(selected_answer) == question["answer"]:
                score += 1


    # -------------------------
    # CALCULATE RESULT
    # -------------------------

    total = len(questions)

    percentage = (score / total) * 100

    grade = calculate_grade(percentage)

    submitted_at = datetime.now().strftime(
        "%Y-%m-%d %H:%M:%S"
    )


    # -------------------------
    # SAVE RESULT
    # -------------------------

    conn = get_db()

    conn.execute("""
        INSERT INTO results
        (
            student1_id,
            student1_name,
            student2_id,
            student2_name,
            score,
            total,
            percentage,
            grade,
            submitted_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        student1_id,
        student1_name,
        student2_id,
        student2_name,
        score,
        total,
        percentage,
        grade,
        submitted_at
    ))

    conn.commit()
    conn.close()


    # -------------------------
    # SHOW RESULT
    # -------------------------

    return render_template(
        "result.html",

        questions=questions,
        answers=answers,

        score=score,
        total=total,
        percentage=percentage,
        grade=grade,

        student1_name=student1_name,
        student2_name=student2_name
    )
    # -------------------------
    # EXAM SUBMISSION
    # -------------------------

    student1_id = request.form.get("student1_id")
    student1_name = request.form.get("student1_name")

    student2_id = request.form.get("student2_id")
    student2_name = request.form.get("student2_name")

    score = 0

    answers = {}

    for question in questions:

        question_id = str(question["id"])

        selected_answer = request.form.get(
            f"question_{question_id}"
        )

        answers[question_id] = selected_answer

        if selected_answer is not None:

            if int(selected_answer) == question["answer"]:
                score += 1

    total = len(questions)

    percentage = (score / total) * 100

    grade = calculate_grade(percentage)

    submitted_at = datetime.now().strftime(
        "%Y-%m-%d %H:%M:%S"
    )

    conn = get_db()

    conn.execute("""
        INSERT INTO results
        (
            student1_id,
            student1_name,
            student2_id,
            student2_name,
            score,
            total,
            percentage,
            grade,
            submitted_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        student1_id,
        student1_name,
        student2_id,
        student2_name,
        score,
        total,
        percentage,
        grade,
        submitted_at
    ))

    conn.commit()
    conn.close()

    return render_template(
        "result.html",

        questions=questions,
        answers=answers,

        score=score,
        total=total,
        percentage=percentage,
        grade=grade,

        student1_name=student1_name,
        student2_name=student2_name
    )




@app.route("/admin-login", methods=["GET", "POST"])
def admin_login():

    if request.method == "POST":

        password = request.form.get("password")

        if password == ADMIN_PASSWORD:

            session["admin_logged_in"] = True

            return redirect(url_for("admin"))

        return render_template(
            "admin_login.html",
            error="Incorrect password."
        )

    return render_template("admin_login.html")
# -----------------------------
# ADMIN RESULTS
# -----------------------------

@app.route("/admin")
def admin():

    if not session.get("admin_logged_in"):
        return redirect(url_for("admin_login"))

    conn = get_db()

    results = conn.execute("""
        SELECT *
        FROM results
        ORDER BY score DESC, submitted_at ASC
    """).fetchall()

    conn.close()

    ranked_results = []

    previous_score = None
    current_rank = 0

    for index, result in enumerate(results):

        if result["score"] != previous_score:

            current_rank = index + 1
            previous_score = result["score"]

        ranked_results.append({
            "rank": current_rank,

            "student1_id": result["student1_id"],
            "student1_name": result["student1_name"],

            "student2_id": result["student2_id"],
            "student2_name": result["student2_name"],

            "score": result["score"],
            "total": result["total"],
            "percentage": result["percentage"],
            "grade": result["grade"],
            "submitted_at": result["submitted_at"]
        })

    total_students = len(ranked_results)

    average_score = 0

    if total_students > 0:

        average_score = sum(
            result["percentage"]
            for result in ranked_results
        ) / total_students

    return render_template(
        "admin.html",
        results=ranked_results,
        total_students=total_students,
        average_score=average_score
    )


@app.route("/admin-logout")
def admin_logout():

    session.pop("admin_logged_in", None)

    return redirect(url_for("admin_login"))

@app.route("/export-results")
def export_results():

    # Protect the export
    if not session.get("admin_logged_in"):
        return redirect(url_for("admin_login"))

    conn = get_db()

    results = conn.execute("""
        SELECT *
        FROM results
        ORDER BY score DESC, submitted_at ASC
    """).fetchall()

    conn.close()

    # Create workbook
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "CAT Results"

    # Title
    worksheet.merge_cells("A1:I1")

    worksheet["A1"] = "COMPUTER APPLICATIONS CAT RESULTS"

    worksheet["A1"].font = Font(
        bold=True,
        size=16
    )

    worksheet["A1"].alignment = Alignment(
        horizontal="center"
    )

    # Subtitle
    worksheet.merge_cells("A2:I2")

    worksheet["A2"] = "Microsoft Word & Basic Computer Knowledge"

    worksheet["A2"].font = Font(
        italic=True,
        size=11
    )

    worksheet["A2"].alignment = Alignment(
        horizontal="center"
    )

    # Headers
    headers = [
        "Position",
        "Student 1 ID",
        "Student 1 Name",
        "Student 2 ID",
        "Student 2 Name",
        "Score",
        "Percentage",
        "Grade",
        "Submitted At"
    ]

    header_row = 4

    for column, header in enumerate(headers, start=1):

        cell = worksheet.cell(
            row=header_row,
            column=column,
            value=header
        )

        cell.font = Font(bold=True)

        cell.alignment = Alignment(
            horizontal="center"
        )

    # Ranking
    previous_score = None
    current_rank = 0

    # Data starts on row 5
    for index, result in enumerate(results):

        score = result["score"]

        if score != previous_score:
            current_rank = index + 1
            previous_score = score

        row = index + 5

        data = [
            current_rank,
            result["student1_id"],
            result["student1_name"],
            result["student2_id"],
            result["student2_name"],
            f"{result['score']}/{result['total']}",
            result["percentage"] / 100,
            result["grade"],
            result["submitted_at"]
        ]

        for column, value in enumerate(data, start=1):

            cell = worksheet.cell(
                row=row,
                column=column,
                value=value
            )

            cell.alignment = Alignment(
                vertical="center"
            )

        # Format percentage
        worksheet.cell(
            row=row,
            column=7
        ).number_format = "0.0%"

    # Borders
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    for row in worksheet.iter_rows(
        min_row=4,
        max_row=worksheet.max_row,
        min_col=1,
        max_col=9
    ):

        for cell in row:
            cell.border = thin_border

    # Column widths
    widths = {
        "A": 12,
        "B": 18,
        "C": 25,
        "D": 18,
        "E": 25,
        "F": 12,
        "G": 15,
        "H": 10,
        "I": 22
    }

    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width

    # Freeze headers
    worksheet.freeze_panes = "A5"

    # Create file in memory
    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="Computer_CAT_Results.xlsx",
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )



# -----------------------------
# START APPLICATION
# -----------------------------

if __name__ == "__main__":

    init_db()

    app.run(
        host="0.0.0.0",
        port=5000,
        debug=True
    )

